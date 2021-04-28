# 全流程第三步
# 用于生成word 表格，填写贝叶斯分析数据
from docx import Document
import numpy as np
import scipy.io as io
import os

'''
#表格示例
大小7*5
参数	     theat_t	****	先验分布	****
Lgrad	 0.403		****   U[0.200,0.600]	****
推断次数	 map	  map与theat_t的误差%	 95%置信区间	后验分布与先验分布的区间宽度百分比%
1	0.405	0.42	[0.358,0.446]	22.00
2	0.403	0.06	[0.383,0.427]	11.00
3	0.402	0.14	[0.387,0.420]	8.25
4	0.402	0.20	[0.390,0.417]	6.75
'''
# 创建文档对象函数
def produce_docx(path,data,param):
    document = Document()

    # 创建7行5列表格
    table = document.add_table(rows=7, cols=5,style='Table Grid')

    table.cell(0, 0).text = '参数'
    table.cell(0, 1).text = 'theat_t'
    table.cell(0, 3).text = '先验分布'

    table.cell(1, 0).text = param
    table.cell(1, 1).text = str.format("{:.3f}",data[0,0])
    table.cell(1, 3).text = 'U['+ str.format("{:.3f}",data[0,2]) + ','+ str.format("{:.3f}",data[0,3]) + ']'

    table.cell(2, 0).text = '推断次数'
    table.cell(2, 1).text = 'map'
    table.cell(2, 2).text = 'map与theat_t的误差%'
    table.cell(2, 3).text = '95%置信区间'
    table.cell(2, 4).text = '后验分布与先验分布的区间宽度百分比%'

    for j in range(3, 7):
        table.cell(j, 0).text = str(j-2)
        table.cell(j, 1).text = str.format("{:.3f}", data[j - 2, 0])
        table.cell(j, 2).text = str.format("{:.2f}", data[j - 2, 1])
        table.cell(j, 3).text = '['+ str.format("{:.3f}", data[j-2, 2]) + ','+ str.format("{:.3f}", data[j-2, 3]) + ']'
        table.cell(j, 4).text = str.format("{:.2f}", data[j - 2, 4])
        # 保存文档
    document.save(path + '\\' + param+'table.docx')
    return None

params = ['Lgrad','vc','h','m','rou']

main_path = os.getcwd()
time_file_name = 'Time2021_04_28_08_15_12_function_model_bayes_inference'
path = main_path + '\\' + time_file_name

#path = 'C:\\Users\\Wang haitao\PycharmProjects\\bayes_test\\Time2021_04_28_07_28_40_function_model_bayes_inference'
for  param in params:
    data = io.loadmat(path + '\\' + param + '_bayes_analysis_compute.mat')[param+'_c']
    produce_docx(path,data,param)